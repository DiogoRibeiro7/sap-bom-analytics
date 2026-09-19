"""Adapters for common SAP export file formats."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Final

from openpyxl import load_workbook

_TEXT_SUFFIXES: Final[frozenset[str]] = frozenset({".csv", ".txt", ".tsv"})
_EXCEL_SUFFIXES: Final[frozenset[str]] = frozenset({".xlsx"})


@dataclass(frozen=True, slots=True)
class SapExport:
    """Normalized tabular representation of one SAP export."""

    headers: tuple[str, ...]
    rows: tuple[tuple[int, dict[str, str]], ...]
    format_name: str


def _text_delimiter(path: Path, sample: str) -> str:
    """Infer a delimiter from common SAP text export formats."""
    if path.suffix.lower() == ".tsv":
        return "\t"

    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        return ","
    return dialect.delimiter


def _read_text_export(path: Path) -> SapExport:
    """Read CSV, semicolon, pipe, or tab-delimited SAP text exports."""
    text = path.read_text(encoding="utf-8-sig")
    if not text.strip():
        raise ValueError(f"{path.name}: export is empty")

    delimiter = _text_delimiter(path, text[:8192])
    reader = csv.DictReader(text.splitlines(), delimiter=delimiter)
    if reader.fieldnames is None:
        raise ValueError(f"{path.name}: export has no header")

    headers = tuple(reader.fieldnames)
    rows: list[tuple[int, dict[str, str]]] = []
    for row_number, row in enumerate(reader, start=2):
        normalized: dict[str, str] = {}
        for header in headers:
            value = row.get(header)
            normalized[header] = "" if value is None else str(value)
        rows.append((row_number, normalized))

    return SapExport(
        headers=headers,
        rows=tuple(rows),
        format_name=f"text:{repr(delimiter)}",
    )


def _read_xlsx_export(path: Path) -> SapExport:
    """Read the active worksheet from an XLSX SAP export."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        try:
            first_row = next(values)
        except StopIteration as exc:
            raise ValueError(f"{path.name}: workbook is empty") from exc

        headers = tuple("" if value is None else str(value) for value in first_row)
        rows: list[tuple[int, dict[str, str]]] = []
        for row_number, values_row in enumerate(values, start=2):
            row = {
                header: "" if value is None else str(value)
                for header, value in zip(headers, values_row, strict=False)
            }
            rows.append((row_number, row))
    finally:
        workbook.close()

    return SapExport(headers=headers, rows=tuple(rows), format_name="xlsx")


def read_sap_export(path: Path) -> SapExport:
    """Read a supported SAP export into a common tabular representation."""
    if not isinstance(path, Path):
        raise TypeError("path must be a pathlib.Path")
    if not path.is_file():
        raise FileNotFoundError(path)

    suffix = path.suffix.lower()
    if suffix in _TEXT_SUFFIXES:
        return _read_text_export(path)
    if suffix in _EXCEL_SUFFIXES:
        return _read_xlsx_export(path)

    supported = ", ".join(sorted(_TEXT_SUFFIXES | _EXCEL_SUFFIXES))
    raise ValueError(f"Unsupported SAP export format {suffix!r}; expected one of: {supported}")
